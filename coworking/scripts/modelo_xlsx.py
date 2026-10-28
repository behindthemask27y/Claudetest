# -*- coding: utf-8 -*-
"""Modelo financiero editable (openpyxl) del coworking de Santo Tomé 2987.

Lee los supuestos de datos/supuestos.json y escribe en el libro:
  Supuestos · Inversion · Modelo_Conservador · Modelo_Base · Modelo_Favorable · Resumen · Equilibrio · Sensibilidad
Todo con fórmulas que dependen de Supuestos (celdas azules = entradas).
Marco monetario: USD constantes a la fecha de corte; los precios en ARS se pasan con el TC de Supuestos.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

NAVY, GOLD, INK, GRAY = "253C64", "988355", "171717", "6A6A6A"
F_IN = Font(name="Arial", size=10, color="1F4E9E")            # entrada editable (azul)
F_TX = Font(name="Arial", size=10, color="2B2B34")
F_B = Font(name="Arial", size=10, bold=True, color=INK)
F_H = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_T = Font(name="Arial", size=14, bold=True, color=INK)
F_S = Font(name="Arial", size=9, italic=True, color=GRAY)
FILL_H = PatternFill("solid", fgColor=NAVY)
FILL_IN = PatternFill("solid", fgColor="EEF2FA")
FILL_SEC = PatternFill("solid", fgColor="DBDBDB")
FILL_KEY = PatternFill("solid", fgColor="FBF0D9")
THIN = Border(bottom=Side(style="thin", color="D2D2D2"))
ESC = [("Conservador", "D"), ("Base", "E"), ("Favorable", "F")]
NMES = 36
USD = '#,##0;[Red]-#,##0'
USD2 = '#,##0.00;[Red]-#,##0.00'
PCT = '0.0%'


def _hdr(ws, row, vals, widths=None):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill = F_H, FILL_H
        c.alignment = Alignment(vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w


def hoja_supuestos(wb, sup):
    ws = wb.create_sheet("Supuestos")
    ws["A1"] = "Supuestos del modelo (celdas azules = entradas editables)"
    ws["A1"].font = F_T
    ws["A2"] = (f"Marco monetario: USD constantes al {sup['fecha_corte']}. TC de conversión de precios en pesos: ver fila 'tc'. "
                "Sin previsiones de inflación ni devaluación. Cada valor tiene su fuente o se marca como supuesto a validar.")
    ws["A2"].font = F_S
    _hdr(ws, 4, ["Clave", "Supuesto", "Unidad", "Conservador", "Base", "Favorable", "Fuente / justificación", "Tipo"],
         [16, 46, 16, 14, 14, 14, 90, 18])
    R, r = {}, 5
    for it in sup["items"]:
        if it.get("seccion"):
            c = ws.cell(row=r, column=1, value=it["seccion"])
            c.font = F_B
            for col in range(1, 9):
                ws.cell(row=r, column=col).fill = FILL_SEC
            r += 1
            continue
        R[it["key"]] = r
        ws.cell(row=r, column=1, value=it["key"]).font = F_S
        ws.cell(row=r, column=2, value=it["label"]).font = F_TX
        ws.cell(row=r, column=3, value=it.get("unit", "")).font = F_S
        vals = it["vals"] if isinstance(it["vals"], list) else [it["vals"]] * 3
        fmt = {"%": PCT, "USD": USD, "ARS": '#,##0', "n": '#,##0.0'}.get(it.get("fmt", "n"), '#,##0.0')
        for j, v in enumerate(vals):
            c = ws.cell(row=r, column=4 + j, value=v)
            c.font, c.fill, c.number_format = F_IN, FILL_IN, fmt
        ws.cell(row=r, column=7, value=it.get("fuente", "")).font = F_S
        ws.cell(row=r, column=7).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=8, value=it.get("tipo", "")).font = F_S
        r += 1
    ws.freeze_panes = "D5"
    return R


def S(R, key, col):
    """referencia absoluta a un supuesto en la columna del escenario"""
    return f"Supuestos!${col}${R[key]}"


def hoja_inversion(wb, R):
    ws = wb.create_sheet("Inversion")
    ws["A1"] = "Inversión inicial (USD, mes 0)"
    ws["A1"].font = F_T
    ws["A2"] = "Se calcula desde Supuestos. El capital de trabajo no se suma acá: es la necesidad máxima de caja de cada escenario (ver Resumen)."
    ws["A2"].font = F_S
    _hdr(ws, 4, ["Rubro", "Cómo se calcula", "", "Conservador", "Base", "Favorable"], [40, 60, 2, 16, 16, 16])
    lineas = [
        ("Obra de adecuación general", "m² intervenidos × USD/m²", lambda c: f"={S(R,'m2_obra',c)}*{S(R,'usd_m2_obra',c)}"),
        ("Techado del patio", "m² de patio × USD/m²", lambda c: f"={S(R,'m2_patio',c)}*{S(R,'usd_m2_patio',c)}"),
        ("Anulación / relleno de pileta", "suma fija", lambda c: f"={S(R,'pileta_usd',c)}"),
        ("Baño accesible", "suma fija", lambda c: f"={S(R,'bano_acc_usd',c)}"),
        ("Instalación eléctrica y red de datos", "suma fija", lambda c: f"={S(R,'electrica_datos_usd',c)}"),
        ("Climatización", "equipos × USD por equipo instalado", lambda c: f"={S(R,'clima_n',c)}*{S(R,'clima_usd',c)}"),
        ("Mobiliario de puestos", "(puestos flex + fijos + oficinas) × USD por puesto", lambda c: f"=({S(R,'puestos_flex',c)}+{S(R,'puestos_fijos',c)}+{S(R,'oficinas_priv',c)}*{S(R,'puestos_por_oficina',c)})*{S(R,'mueble_puesto_usd',c)}"),
        ("Mobiliario común (salas, living, lockers)", "suma fija", lambda c: f"={S(R,'mueble_comun_usd',c)}"),
        ("Cabinas para llamadas", "cabinas × USD", lambda c: f"={S(R,'cabinas_n',c)}*{S(R,'cabina_usd',c)}"),
        ("Café y cocina (equipamiento)", "suma fija", lambda c: f"={S(R,'cafe_equipo_usd',c)}"),
        ("Conectividad, control de acceso y cámaras", "suma fija", lambda c: f"={S(R,'it_seguridad_usd',c)}"),
    ]
    r = 5
    for nom, como, f in lineas:
        ws.cell(row=r, column=1, value=nom).font = F_TX
        ws.cell(row=r, column=2, value=como).font = F_S
        for _, c in ESC:
            col = {"D": 4, "E": 5, "F": 6}[c]
            ws.cell(row=r, column=col, value=f(c)).number_format = USD
        r += 1
    sub_ini, sub_fin = 5, r - 1
    ws.cell(row=r, column=1, value="Subtotal obra y equipamiento").font = F_B
    for c in "DEF":
        ws[f"{c}{r}"] = f"=SUM({c}{sub_ini}:{c}{sub_fin})"
        ws[f"{c}{r}"].number_format, ws[f"{c}{r}"].font = USD, F_B
    r_sub = r
    r += 1
    extra = [
        ("Honorarios profesionales (proyecto, dirección, planos)", "% sobre obra (filas de obra e instalaciones)",
         lambda c: f"=SUM({c}5:{c}9)*{S(R,'honorarios_pct',c)}"),
        ("Habilitación, planos de uso, incendio y trámites", "suma fija", lambda c: f"={S(R,'habilitacion_usd',c)}"),
        ("Marketing de apertura y preventa", "suma fija", lambda c: f"={S(R,'apertura_mkt_usd',c)}"),
    ]
    for nom, como, f in extra:
        ws.cell(row=r, column=1, value=nom).font = F_TX
        ws.cell(row=r, column=2, value=como).font = F_S
        for c in "DEF":
            ws[f"{c}{r}"] = f(c)
            ws[f"{c}{r}"].number_format = USD
        r += 1
    ws.cell(row=r, column=1, value="Contingencias").font = F_TX
    ws.cell(row=r, column=2, value="% sobre todo lo anterior").font = F_S
    for c in "DEF":
        ws[f"{c}{r}"] = f"=SUM({c}{r_sub}:{c}{r-1})*{S(R,'contingencia_pct',c)}"
        ws[f"{c}{r}"].number_format = USD
    r += 1
    ws.cell(row=r, column=1, value="INVERSIÓN INICIAL TOTAL (sin capital de trabajo)").font = F_B
    for c in "DEF":
        ws[f"{c}{r}"] = f"=SUM({c}{r_sub}:{c}{r-1})"
        ws[f"{c}{r}"].number_format, ws[f"{c}{r}"].font, ws[f"{c}{r}"].fill = USD, F_B, FILL_KEY
    ws.freeze_panes = "D5"
    return r   # fila del total


def hoja_modelo(wb, R, nombre, c, fila_inv):
    """36 meses. Columna C = mes 0 (inversión); D..AM = meses 1..36."""
    ws = wb.create_sheet(f"Modelo_{nombre}")
    ws["A1"] = f"Modelo mensual a 36 meses · escenario {nombre} (USD constantes)"
    ws["A1"].font = F_T
    ws["A2"] = "Miembros = fin de mes. Ingresos sin IVA. Flujo de caja con la casa propia (sin alquiler pagado); abajo, el resultado económico restando el alquiler que se deja de cobrar."
    ws["A2"].font = F_S
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 52
    ws.cell(row=3, column=2, value="Mes").font = F_H
    ws.cell(row=3, column=2).fill = FILL_H
    for m in range(0, NMES + 1):
        cc = ws.cell(row=3, column=3 + m, value=m)
        cc.font, cc.fill, cc.alignment = F_H, FILL_H, Alignment(horizontal="center")
        ws.column_dimensions[L(3 + m)].width = 11
    s = lambda k: S(R, k, c)
    dprom = (f"({s('mix_p1')}*{s('dias_p1')}+{s('mix_p2')}*{s('dias_p2')}+{s('mix_p3')}*{s('dias_p3')}+{s('mix_p4')}*{s('dias_p4')})")
    pprom = (f"(({s('mix_p1')}*{s('precio_p1')}+{s('mix_p2')}*{s('precio_p2')}+{s('mix_p3')}*{s('precio_p3')}+{s('mix_p4')}*{s('precio_p4')})*{s('precio_mult')})")
    neto = f"*(1-{s('descuento')})/(1+{s('iva')}*{s('ri')})/{s('tc')}"   # ARS con IVA -> USD sin IVA
    cap_flex = f"({s('puestos_flex')}*{s('dias_habiles')}/({dprom}*{s('uso_flex')}*{s('factor_pico')}))"

    filas = {}
    defs = [
        # clave, etiqueta, fórmula(m, col, prev_col) o None (sección), formato
        ("sec1", "MIEMBROS Y OCUPACIÓN", None, None),
        ("cap", "Tope de miembros flex que entran en los puestos (pico)", lambda m, x, p: f"={cap_flex}", '0.0'),
        ("altas", "Altas de miembros flex en el mes", lambda m, x, p: f"=IF({x}$3<=6,{s('altas_f1')},IF({x}$3<=12,{s('altas_f2')},{s('altas_f3')}))", '0.0'),
        ("flex", "Miembros flex (fin de mes)", lambda m, x, p: (f"=MIN({x}{{cap}},{s('miembros_ini')}*(1-{s('churn')})+{x}{{altas}})" if m == 1 else f"=MIN({x}{{cap}},{p}{{flex}}*(1-{s('churn')})+{x}{{altas}})"), '0.0'),
        ("bajas", "Bajas estimadas de flex en el mes", lambda m, x, p: (f"={s('miembros_ini')}*{s('churn')}" if m == 1 else f"={p}{{flex}}*{s('churn')}"), '0.0'),
        ("fijos", "Puestos fijos ocupados", lambda m, x, p: (f"=MIN({s('puestos_fijos')},{s('altas_fijo')})" if m == 1 else f"=MIN({s('puestos_fijos')},{p}{{fijos}}*(1-{s('churn_fijo')})+{s('altas_fijo')})"), '0.0'),
        ("ofic", "Oficinas privadas ocupadas", lambda m, x, p: f"={s('oficinas_priv')}*MIN(1,{x}$3/{s('ramp_of')})", '0.0'),
        ("pdias", "Puestos-día usados por miembros flex", lambda m, x, p: f"={x}{{flex}}*{dprom}*{s('uso_flex')}", '#,##0'),
        ("occ", "Ocupación promedio de puestos flex (días hábiles)", lambda m, x, p: f"={x}{{pdias}}/({s('puestos_flex')}*{s('dias_habiles')})", PCT),
        ("pico", "Puestos flex simultáneos en el pico (estimado)", lambda m, x, p: f"={x}{{pdias}}/{s('dias_habiles')}*{s('factor_pico')}", '0.0'),
        ("sec2", "INGRESOS (USD sin IVA)", None, None),
        ("i_flex", "Membresías flex (días por mes / ilimitado)", lambda m, x, p: f"={x}{{flex}}*{pprom}{neto}", USD),
        ("i_fijo", "Puestos fijos", lambda m, x, p: f"={x}{{fijos}}*{s('precio_fijo')}*{s('precio_mult')}{neto}", USD),
        ("i_ofic", "Oficinas privadas (complemento)", lambda m, x, p: f"={x}{{ofic}}*{s('precio_oficina')}*{s('precio_mult')}{neto}", USD),
        ("i_pase", "Pases diarios", lambda m, x, p: f"={s('pases_mes')}*MIN(1,{x}$3/{s('ramp_pases')})*{s('precio_pase')}*{s('precio_mult')}/(1+{s('iva')}*{s('ri')})/{s('tc')}", USD),
        ("i_sala", "Salas por hora a externos", lambda m, x, p: f"={s('horas_sala_mes')}*MIN(1,{x}$3/{s('ramp_pases')})*{s('precio_sala_h')}*{s('precio_mult')}/(1+{s('iva')}*{s('ri')})/{s('tc')}", USD),
        ("i_otros", "Otros (domicilio comercial, eventos)", lambda m, x, p: f"={s('otros_ingresos_usd')}*MIN(1,{x}$3/{s('ramp_pases')})", USD),
        ("ing", "INGRESOS TOTALES", lambda m, x, p: f"=SUM({x}{{i_flex}}:{x}{{i_otros}})", USD),
        ("sec3", "COSTOS VARIABLES", None, None),
        ("v_cafe", "Café, té, agua y consumibles", lambda m, x, p: f"=({x}{{pdias}}+{x}{{fijos}}*{s('dias_habiles')}*{s('uso_fijo')}+({x}{{ofic}}*{s('puestos_por_oficina')}*{s('dias_habiles')}*{s('uso_fijo')})+{s('pases_mes')}*MIN(1,{x}$3/{s('ramp_pases')}))*{s('costo_puesto_dia_usd')}", USD),
        ("v_com", "Comisiones de cobro", lambda m, x, p: f"={x}{{ing}}*(1+{s('iva')}*{s('ri')})*{s('comision_pct')}", USD),
        ("v_iibb", "Ingresos Brutos", lambda m, x, p: f"={x}{{ing}}*{s('iibb_pct')}", USD),
        ("v_idc", "Impuesto a los débitos y créditos (aprox.)", lambda m, x, p: f"={x}{{ing}}*(1+{s('iva')}*{s('ri')})*{s('idc_pct')}", USD),
        ("var", "Total costos variables", lambda m, x, p: f"=SUM({x}{{v_cafe}}:{x}{{v_idc}})", USD),
        ("mc", "MARGEN DE CONTRIBUCIÓN", lambda m, x, p: f"={x}{{ing}}-{x}{{var}}", USD),
        ("sec4", "COSTOS FIJOS MENSUALES", None, None),
        ("f_pers", "Personal (recepción / comunidad), con cargas", lambda m, x, p: f"={s('personal_fte')}*{s('costo_fte_usd')}", USD),
        ("f_limp", "Limpieza", lambda m, x, p: f"={s('limpieza_usd')}", USD),
        ("f_net", "Internet principal + respaldo", lambda m, x, p: f"={s('internet_usd')}", USD),
        ("f_serv", "Electricidad, gas y agua", lambda m, x, p: f"={s('servicios_usd')}", USD),
        ("f_seg", "Seguros", lambda m, x, p: f"={s('seguros_usd')}", USD),
        ("f_soft", "Software de gestión y reservas", lambda m, x, p: f"={s('software_usd')}", USD),
        ("f_adm", "Contador y administración", lambda m, x, p: f"={s('admin_usd')}", USD),
        ("f_mant", "Mantenimiento y reposición", lambda m, x, p: f"=Inversion!{c}{fila_inv}*{s('mant_pct_anual')}/12", USD),
        ("f_mkt", "Publicidad y comunidad (eventos, desayunos)", lambda m, x, p: f"={s('marketing_usd')}+{s('comunidad_usd')}", USD),
        ("f_otros", "Otros e imprevistos", lambda m, x, p: f"=SUM({x}{{f_pers}}:{x}{{f_mkt}})*{s('otros_pct')}", USD),
        ("fij", "Total costos fijos", lambda m, x, p: f"=SUM({x}{{f_pers}}:{x}{{f_otros}})", USD),
        ("sec5", "RESULTADO Y CAJA", None, None),
        ("ebitda", "RESULTADO OPERATIVO (antes de Ganancias)", lambda m, x, p: f"={x}{{mc}}-{x}{{fij}}", USD),
        ("inv", "Inversión inicial", None, USD),
        ("fc", "Flujo de caja del proyecto (casa propia, sin alquiler pagado)", lambda m, x, p: f"={x}{{ebitda}}+{x}{{inv}}", USD),
        ("acum", "Flujo acumulado", lambda m, x, p: f"={p}{{acum}}+{x}{{fc}}", USD),
        ("op", "Alquiler de la casa que se deja de cobrar (costo de oportunidad)", lambda m, x, p: f"=-{s('alquiler_casa_usd')}", USD),
        ("eco", "Resultado económico frente a alquilar la casa", lambda m, x, p: f"={x}{{fc}}+{x}{{op}}", USD),
        ("eacum", "Resultado económico acumulado", lambda m, x, p: f"={p}{{eacum}}+{x}{{eco}}", USD),
    ]
    r = 4
    for k, lab, f, fmt in defs:
        filas[k] = r
        r += 1
    for k, lab, f, fmt in defs:
        rr = filas[k]
        cl = ws.cell(row=rr, column=2, value=lab)
        if f is None and k.startswith("sec"):
            cl.font = F_B
            for col in range(2, 4 + NMES):
                ws.cell(row=rr, column=col).fill = FILL_SEC
            continue
        cl.font = F_B if k in ("ing", "mc", "fij", "ebitda", "fc", "eco") else F_TX
        # mes 0
        c0 = ws.cell(row=rr, column=3)
        if k == "inv":
            c0.value = f"=-Inversion!{c}{fila_inv}"
        elif k in ("fc",):
            c0.value = f"=C{filas['inv']}"
        elif k == "acum":
            c0.value = f"=C{filas['fc']}"
        elif k == "op":
            c0.value = f"=-{s('alquiler_casa_usd')}*{s('meses_obra')}"
        elif k == "eco":
            c0.value = f"=C{filas['fc']}+C{filas['op']}"
        elif k == "eacum":
            c0.value = f"=C{filas['eco']}"
        if c0.value is not None:
            c0.number_format = fmt or USD
        for m in range(1, NMES + 1):
            col = 3 + m
            x, p = L(col), L(col - 1)
            if k == "inv":
                ws.cell(row=rr, column=col, value=0).number_format = USD
                continue
            txt = f(m, x, p)
            for kk, vv in filas.items():
                txt = txt.replace("{" + kk + "}", str(vv))
            cc = ws.cell(row=rr, column=col, value=txt)
            cc.number_format = fmt or USD
            if k in ("ebitda", "fc", "eco"):
                cc.font = F_B
    ws.freeze_panes = "C4"
    return filas


def hoja_resumen(wb, R, filas, fila_inv):
    ws = wb.create_sheet("Resumen", 0)
    ws["A1"] = "Resumen por escenario"
    ws["A1"].font = F_T
    ws["A2"] = "Todo se recalcula al cambiar Supuestos. USD constantes; resultado antes del impuesto a las Ganancias."
    ws["A2"].font = F_S
    _hdr(ws, 4, ["Indicador", "Conservador", "Base", "Favorable", "Cómo se lee"], [58, 16, 16, 16, 80])
    last = L(3 + NMES)
    def M(n):
        return f"Modelo_{n}"
    rows = [
        ("Inversión inicial (sin capital de trabajo)", lambda n, c: f"=Inversion!{c}{fila_inv}", USD, "Obra, equipamiento, honorarios, trámites, apertura y contingencias."),
        ("Necesidad máxima de caja (inversión + pérdidas de arranque)", lambda n, c: f"=-MIN({M(n)}!C{filas['acum']}:{last}{filas['acum']})", USD, "Lo que hay que tener disponible en el peor momento del flujo acumulado."),
        ("Miembros flex en el mes 12", lambda n, c: f"={M(n)}!O{filas['flex']}", '0', ""),
        ("Miembros flex en el mes 36", lambda n, c: f"={M(n)}!{last}{filas['flex']}", '0', "Limitado por el tope de puestos en el pico."),
        ("Ocupación promedio de puestos flex, mes 36", lambda n, c: f"={M(n)}!{last}{filas['occ']}", PCT, "Puestos-día usados / puestos-día disponibles en días hábiles."),
        ("Ingresos del mes 36 (USD sin IVA)", lambda n, c: f"={M(n)}!{last}{filas['ing']}", USD, ""),
        ("Resultado operativo del mes 36", lambda n, c: f"={M(n)}!{last}{filas['ebitda']}", USD, ""),
        ("Resultado operativo acumulado de los 36 meses", lambda n, c: f"=SUM({M(n)}!D{filas['ebitda']}:{last}{filas['ebitda']})", USD, ""),
        ("Margen operativo del año 3", lambda n, c: f"=IFERROR(SUM({M(n)}!AB{filas['ebitda']}:{last}{filas['ebitda']})/SUM({M(n)}!AB{filas['ing']}:{last}{filas['ing']}),0)", PCT, "Resultado operativo / ingresos, meses 25 a 36."),
        ("Mes de recupero de la inversión (caja, casa propia)", lambda n, c: f"=IFERROR(MATCH(TRUE,INDEX({M(n)}!C{filas['acum']}:{last}{filas['acum']}>=0,0),0)-1,\"no recupera en 36 meses\")", '0', "Primer mes con flujo acumulado ≥ 0."),
        ("Mes de recupero frente a alquilar la casa", lambda n, c: f"=IFERROR(MATCH(TRUE,INDEX({M(n)}!C{filas['eacum']}:{last}{filas['eacum']}>=0,0),0)-1,\"no recupera en 36 meses\")", '0', "Igual, restando cada mes el alquiler que se deja de cobrar."),
        ("TIR anual a 36 meses, sin valor residual (caja)", lambda n, c: f"=IFERROR((1+IRR({M(n)}!C{filas['fc']}:{last}{filas['fc']}))^12-1,\"no calculable\")", PCT, "Conservadora: no suma lo que vale el negocio ni el equipamiento al mes 36."),
        ("VAN a 36 meses con la tasa de Supuestos (caja)", lambda n, c: f"=NPV((1+{S(R,'tasa_desc_anual',c)})^(1/12)-1,{M(n)}!D{filas['fc']}:{last}{filas['fc']})+{M(n)}!C{filas['fc']}+{S(R,'valor_residual_usd',c)}/(1+{S(R,'tasa_desc_anual',c)})^3", USD, "La tasa la define el inversor; el valor que viene es un supuesto."),
        ("Alquiler de mercado de la casa que se deja de cobrar (por mes)", lambda n, c: f"={S(R,'alquiler_casa_usd',c)}", USD, "Costo de oportunidad de usar la casa propia."),
        ("Resultado operativo del año 3 menos el alquiler que se deja de cobrar (por mes)", lambda n, c: f"=AVERAGE({M(n)}!AB{filas['ebitda']}:{last}{filas['ebitda']})-{S(R,'alquiler_casa_usd',c)}", USD, "Si es negativo, alquilar la casa rinde más que el coworking (antes de impuestos)."),
    ]
    r = 5
    for lab, f, fmt, nota in rows:
        ws.cell(row=r, column=1, value=lab).font = F_TX
        for j, (n, c) in enumerate(ESC):
            cc = ws.cell(row=r, column=2 + j, value=f(n, c))
            cc.number_format, cc.font = fmt, F_B
        ws.cell(row=r, column=5, value=nota).font = F_S
        r += 1
    ws.freeze_panes = "B5"


def hoja_equilibrio(wb, R, filas):
    ws = wb.create_sheet("Equilibrio")
    ws["A1"] = "Punto de equilibrio en régimen (mes tipo, con los supuestos de cada escenario)"
    ws["A1"].font = F_T
    ws["A2"] = "¿Cuántos miembros flex hacen falta para cubrir los costos fijos, con los puestos fijos, oficinas y extras del mes 36? Comparado con el tope de puestos."
    ws["A2"].font = F_S
    _hdr(ws, 4, ["Concepto", "Conservador", "Base", "Favorable", "Fórmula"], [58, 16, 16, 16, 70])
    last = L(3 + NMES)
    def s(k, c):
        return S(R, k, c)
    def dprom(c):
        return f"({s('mix_p1',c)}*{s('dias_p1',c)}+{s('mix_p2',c)}*{s('dias_p2',c)}+{s('mix_p3',c)}*{s('dias_p3',c)}+{s('mix_p4',c)}*{s('dias_p4',c)})"
    def pnet(c):
        return (f"(({s('mix_p1',c)}*{s('precio_p1',c)}+{s('mix_p2',c)}*{s('precio_p2',c)}+{s('mix_p3',c)}*{s('precio_p3',c)}+{s('mix_p4',c)}*{s('precio_p4',c)})"
                f"*{s('precio_mult',c)}*(1-{s('descuento',c)})/(1+{s('iva',c)}*{s('ri',c)})/{s('tc',c)})")
    def pct_var(c):
        return f"((1+{s('iva',c)}*{s('ri',c)})*({s('comision_pct',c)}+{s('idc_pct',c)})+{s('iibb_pct',c)})"
    rows = [
        ("Precio neto promedio por miembro flex (USD/mes, sin IVA)", lambda n, c: f"={pnet(c)}", USD2, "mezcla de planes × precio × (1 − descuento) / IVA / TC"),
        ("Días promedio incluidos por miembro flex", lambda n, c: f"={dprom(c)}", '0.0', "mezcla × días de cada plan"),
        ("Costo variable por miembro flex (USD/mes)", lambda n, c: f"={dprom(c)}*{s('uso_flex',c)}*{s('costo_puesto_dia_usd',c)}+{pnet(c)}*{pct_var(c)}", USD2, "consumo + comisiones + impuestos sobre ventas"),
        ("Contribución por miembro flex (USD/mes)", lambda n, c: f"={L(2+['D','E','F'].index(c))}5-{L(2+['D','E','F'].index(c))}7", USD2, "precio neto − costo variable"),
        ("Costos fijos del mes 36", lambda n, c: f"=Modelo_{n}!{last}{filas['fij']}", USD, ""),
        ("Contribución de fijos, oficinas y extras del mes 36", lambda n, c: f"=Modelo_{n}!{last}{filas['mc']}-Modelo_{n}!{last}{filas['flex']}*{L(2+['D','E','F'].index(c))}8", USD, "margen de contribución total − el de los miembros flex"),
        ("MIEMBROS FLEX NECESARIOS PARA EL EQUILIBRIO", lambda n, c: f"=MAX(0,({L(2+['D','E','F'].index(c))}9-{L(2+['D','E','F'].index(c))}10)/{L(2+['D','E','F'].index(c))}8)", '0.0', "(fijos − otras contribuciones) / contribución por miembro"),
        ("Tope de miembros flex (puestos en el pico)", lambda n, c: f"=Modelo_{n}!{last}{filas['cap']}", '0.0', ""),
        ("Equilibrio como % del tope", lambda n, c: f"={L(2+['D','E','F'].index(c))}11/{L(2+['D','E','F'].index(c))}12", PCT, "más de 80 %: el negocio depende de estar casi lleno"),
    ]
    r = 5
    for lab, f, fmt, nota in rows:
        ws.cell(row=r, column=1, value=lab).font = F_B if "NECESARIOS" in lab else F_TX
        for j, (n, c) in enumerate(ESC):
            cc = ws.cell(row=r, column=2 + j, value=f(n, c))
            cc.number_format = fmt
            if "NECESARIOS" in lab:
                cc.font, cc.fill = F_B, FILL_KEY
        ws.cell(row=r, column=5, value=nota).font = F_S
        r += 1


def hoja_sensibilidad(wb, R, filas):
    """Resultado operativo de un mes en régimen (escenario Base) ante ocupación y precio. Fórmulas vivas."""
    ws = wb.create_sheet("Sensibilidad")
    ws["A1"] = "Sensibilidad del resultado operativo mensual en régimen · escenario Base"
    ws["A1"].font = F_T
    ws["A2"] = ("Miembros flex = % del tope de puestos. Precio = variación sobre los precios de Supuestos. "
                "Fijos, oficinas y extras como en el mes 36 del escenario Base. USD por mes, antes de Ganancias.")
    ws["A2"].font = F_S
    ws.column_dimensions["A"].width = 30
    e = "Equilibrio"
    ws["A4"] = "Miembros flex (% del tope) ↓ / Precio →"
    ws["A4"].font = F_B
    precios = [-0.2, -0.1, 0, 0.1, 0.2]
    occs = [0.4, 0.5, 0.6, 0.7, 0.8, 0.9, 1.0]
    for j, dp in enumerate(precios):
        c = ws.cell(row=4, column=2 + j, value=dp)
        c.number_format, c.font, c.fill = '+0%;-0%;0%', F_H, FILL_H
        ws.column_dimensions[L(2 + j)].width = 14
    pv = f"((1+Supuestos!$E${R['iva']}*Supuestos!$E${R['ri']})*(Supuestos!$E${R['comision_pct']}+Supuestos!$E${R['idc_pct']})+Supuestos!$E${R['iibb_pct']})"
    for i, oc in enumerate(occs):
        rr = 5 + i
        c = ws.cell(row=rr, column=1, value=oc)
        c.number_format, c.font = '0%', F_B
        for j in range(len(precios)):
            col = L(2 + j)
            # miembros × (precio neto × (1+Δ) × (1 − %var) − consumo) + otras contribuciones × (1+Δ) − fijos
            f = (f"=$A{rr}*{e}!C12*({e}!C5*(1+{col}$4)*(1-{pv})-({e}!C7-{e}!C5*{pv}))"
                 f"+{e}!C10*(1+{col}$4)-{e}!C9")
            cc = ws.cell(row=rr, column=2 + j, value=f)
            cc.number_format = USD
    ws["A13"] = "Lectura: en rojo, el mes da pérdida. La celda de 0 % de precio y la ocupación que alcanza el escenario Base es la del Resumen."
    ws["A13"].font = F_S
    # tornado simple sobre el mes 36 del escenario Base (variación ±20 %), vivo
    ws["A15"] = "Qué mueve más el resultado (mes en régimen, escenario Base, ±20 %)"
    ws["A15"].font = F_B
    _hdr_row = ["Variable", "−20 %", "+20 %", "Rango"]
    for j, v in enumerate(_hdr_row):
        c = ws.cell(row=16, column=1 + j, value=v)
        c.font, c.fill = F_H, FILL_H
    base = f"({e}!C12*{e}!C11/{e}!C11)"  # placeholder no usado
    mb = f"Modelo_Base!{L(3+NMES)}{filas['flex']}"
    tor = [
        ("Precio de las membresías", f"={mb}*{e}!C8+{e}!C10-{e}!C9+{mb}*{e}!C5*(1-{pv})*(-0.2)",
         f"={mb}*{e}!C8+{e}!C10-{e}!C9+{mb}*{e}!C5*(1-{pv})*(0.2)"),
        ("Cantidad de miembros flex", f"={mb}*0.8*{e}!C8+{e}!C10-{e}!C9", f"=MIN({mb}*1.2,{e}!C12)*{e}!C8+{e}!C10-{e}!C9"),
        ("Costos fijos", f"={mb}*{e}!C8+{e}!C10-{e}!C9*0.8", f"={mb}*{e}!C8+{e}!C10-{e}!C9*1.2"),
        ("Personal (sueldos y cargas)", f"={mb}*{e}!C8+{e}!C10-{e}!C9+Modelo_Base!{L(3+NMES)}{filas['f_pers']}*0.2",
         f"={mb}*{e}!C8+{e}!C10-{e}!C9-Modelo_Base!{L(3+NMES)}{filas['f_pers']}*0.2"),
    ]
    for i, (nom, a, b) in enumerate(tor):
        rr = 17 + i
        ws.cell(row=rr, column=1, value=nom).font = F_TX
        ws.cell(row=rr, column=2, value=a).number_format = USD
        ws.cell(row=rr, column=3, value=b).number_format = USD
        ws.cell(row=rr, column=4, value=f"=ABS(C{rr}-B{rr})").number_format = USD
    ws.cell(row=22, column=1, value="Referencia: resultado del mes 36, escenario Base").font = F_S
    ws.cell(row=22, column=2, value=f"=Modelo_Base!{L(3+NMES)}{filas['ebitda']}").number_format = USD


def armar_modelo(wb, sup):
    R = hoja_supuestos(wb, sup)
    fila_inv = hoja_inversion(wb, R)
    filas = None
    for n, c in ESC:
        filas = hoja_modelo(wb, R, n, c, fila_inv)
    hoja_resumen(wb, R, filas, fila_inv)
    hoja_equilibrio(wb, R, filas)
    hoja_sensibilidad(wb, R, filas)
    return R, filas, fila_inv
